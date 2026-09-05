from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.errors import AppError
from app.core.time import app_now_naive
from app.crm.models.document import CrmDocument
from app.crm.models.document_template import DocumentType
from app.crm.models.inspection import CrmInspectionExport, CrmInspectionSession, InspectionSessionStatus, InspectionViewType
from app.crm.models.order import CrmOrder
from app.crm.schemas.document import DocumentPreviewRead
from app.crm.services.audit_log_service import AuditLogService
from app.crm.services.document_mapping_service import TemplateRenderBundle
from app.crm.services.document_pdf_converter_service import DocumentPdfConverterService
from app.crm.services.document_storage_service import DocumentStorageService
from app.crm.services.document_template_service import DocumentTemplateService
from app.crm.services.inspection_canvas_renderer_service import InspectionCanvasRendererService
from app.crm.services.inspection_service import InspectionService


VIEW_ORDER = [
    InspectionViewType.FRONT,
    InspectionViewType.REAR,
    InspectionViewType.LEFT,
    InspectionViewType.RIGHT,
    InspectionViewType.TOP,
    InspectionViewType.INTERIOR,
]

VIEW_LABELS = {
    InspectionViewType.FRONT: "Вид спереди",
    InspectionViewType.REAR: "Вид сзади",
    InspectionViewType.LEFT: "Левый борт",
    InspectionViewType.RIGHT: "Правый борт",
    InspectionViewType.TOP: "Вид сверху",
    InspectionViewType.INTERIOR: "Салон",
}


class InspectionDocumentService:
    def __init__(self, session: Session) -> None:
        self.session = session
        self.audit_logs = AuditLogService(session)
        self.canvas_renderer = InspectionCanvasRendererService()
        self.document_storage = DocumentStorageService()
        self.document_templates = DocumentTemplateService(session)
        self.inspection_service = InspectionService(session)
        self.pdf_converter = DocumentPdfConverterService()

    def _get_export(self, export_id: int) -> CrmInspectionExport:
        export = self.session.get(CrmInspectionExport, export_id)
        if export is None:
            raise AppError(code="not_found", message="Экспорт осмотра не найден", status_code=404)
        return export

    def _get_export_by_document(self, document_id: int) -> CrmInspectionExport:
        export = self.session.scalar(
            select(CrmInspectionExport).where(CrmInspectionExport.document_id == document_id)
        )
        if export is None:
            raise AppError(
                code="not_found",
                message="Для документа не найден snapshot акта осмотра",
                status_code=404,
            )
        return export

    def _get_document(self, document_id: int) -> CrmDocument:
        statement = (
            select(CrmDocument)
            .options(
                selectinload(CrmDocument.order).selectinload(CrmOrder.client),
                selectinload(CrmDocument.order).selectinload(CrmOrder.vehicle),
                selectinload(CrmDocument.template),
            )
            .where(CrmDocument.id == document_id)
        )
        document = self.session.scalar(statement)
        if document is None:
            raise AppError(code="not_found", message="Документ не найден", status_code=404)
        return document

    def _create_document(self, inspection_session: CrmInspectionSession, actor_user_id: int) -> CrmDocument:
        template = self.document_templates.get_by_type(DocumentType.INSPECTION_ACT)
        now = app_now_naive()
        document = CrmDocument(
            order_id=inspection_session.order_id,
            template_id=template.id,
            created_by_user_id=actor_user_id,
            document_type=DocumentType.INSPECTION_ACT,
            document_number=inspection_session.order_id,
            work_started_at=inspection_session.started_at,
            work_completed_at=inspection_session.completed_at,
            created_at=now,
            updated_at=now,
        )
        self.session.add(document)
        self.session.flush()
        return document

    def _render_export_images(self, inspection_session: CrmInspectionSession) -> dict[str, str]:
        export_dir = self.inspection_service.storage.get_exports_dir(inspection_session.id)
        marks = sorted(inspection_session.marks, key=lambda item: (item.sort_order, item.id))
        export_paths: dict[str, str] = {}
        for view_type in VIEW_ORDER:
            view_marks = [
                {
                    "id": mark.id,
                    "geometry_type": mark.geometry_type.value,
                    "geometry_data": deepcopy(mark.geometry_data),
                    "severity": mark.severity,
                }
                for mark in marks
                if mark.view_type == view_type
            ]
            output_path = export_dir / f"{view_type.value}.png"
            self.canvas_renderer.render_view(view_type, view_marks, output_path)
            export_paths[view_type.value] = self.inspection_service.storage.to_relative_path(output_path)

        inspection_session.export_front_image_path = export_paths.get("front")
        inspection_session.export_rear_image_path = export_paths.get("rear")
        inspection_session.export_left_image_path = export_paths.get("left")
        inspection_session.export_right_image_path = export_paths.get("right")
        inspection_session.export_top_image_path = export_paths.get("top")
        inspection_session.export_interior_image_path = export_paths.get("interior")
        return export_paths

    def build_snapshot(self, inspection_session: CrmInspectionSession) -> dict[str, Any]:
        order = inspection_session.order
        client = order.client
        vehicle = order.vehicle
        marks = sorted(inspection_session.marks, key=lambda item: (item.view_type.value, item.sort_order, item.id))
        export_images = self._render_export_images(inspection_session)

        mark_items = []
        for index, mark in enumerate(marks, start=1):
            mark_items.append(
                {
                    "id": mark.id,
                    "number": index,
                    "view_type": mark.view_type.value,
                    "view_label": VIEW_LABELS[mark.view_type],
                    "zone_key": mark.zone_key,
                    "geometry_type": mark.geometry_type.value,
                    "geometry_data": deepcopy(mark.geometry_data),
                    "defect_type": mark.defect_type.value,
                    "severity": mark.severity.value,
                    "status": mark.status.value,
                    "comment": mark.comment,
                    "photos": [
                        {
                            "id": photo.id,
                            "file_path": photo.file_path,
                            "original_name": photo.original_name,
                        }
                        for photo in mark.photos
                    ],
                }
            )

        return {
            "generated_at": app_now_naive().isoformat(),
            "order": {
                "id": order.id,
                "status": order.status,
                "created_at": order.created_at.isoformat(),
            },
            "client": {
                "full_name": client.full_name,
                "phone": client.phone_display,
            },
            "vehicle": {
                "brand": vehicle.brand,
                "model": vehicle.model,
                "plate_number": vehicle.plate_number_display,
                "vin": vehicle.vin,
                "year": vehicle.year,
                "color": vehicle.color,
            },
            "session": {
                "id": inspection_session.id,
                "status": inspection_session.status.value,
                "started_at": inspection_session.started_at.isoformat() if inspection_session.started_at else None,
                "completed_at": inspection_session.completed_at.isoformat() if inspection_session.completed_at else None,
                "confirmed_at": inspection_session.confirmed_at.isoformat() if inspection_session.confirmed_at else None,
                "locked_at": inspection_session.locked_at.isoformat() if inspection_session.locked_at else None,
                "general_comment": inspection_session.general_comment,
                "checklist_json": deepcopy(inspection_session.checklist_json) if inspection_session.checklist_json else None,
                "snapshot_version": inspection_session.snapshot_version,
                "created_by_user_id": inspection_session.created_by_user_id,
                "updated_by_user_id": inspection_session.updated_by_user_id,
                "inspector_name": inspection_session.updated_by.full_name if inspection_session.updated_by else None,
            },
            "marks": mark_items,
            "general_photos": [
                {
                    "id": photo.id,
                    "file_path": photo.file_path,
                    "original_name": photo.original_name,
                }
                for photo in inspection_session.general_photos
            ],
            "export_images": export_images,
        }

    def build_export_payload(self, snapshot: dict[str, Any]) -> dict[str, Any]:
        checklist = snapshot["session"].get("checklist_json") or {}
        checklist_rows = []
        for key, value in checklist.items():
            if isinstance(value, dict):
                checklist_rows.append(
                    {
                        "key": key,
                        "checked": bool(value.get("checked")),
                        "note": value.get("note"),
                    }
                )
            else:
                checklist_rows.append({"key": key, "checked": bool(value), "note": None})

        defects = [
            {
                "number": item["number"],
                "view_label": item["view_label"],
                "defect_type": item["defect_type"],
                "severity": item["severity"],
                "status": item["status"],
                "comment": item["comment"],
                "photos": item["photos"],
            }
            for item in snapshot["marks"]
        ]

        return {
            "document_title": "Акт осмотра автомобиля",
            "order_id": snapshot["order"]["id"],
            "client_name": snapshot["client"]["full_name"],
            "vehicle_display": " ".join(part for part in [snapshot["vehicle"].get("brand"), snapshot["vehicle"].get("model")] if part),
            "checklist_rows": checklist_rows,
            "defects": defects,
            "views": [
                {
                    "view_type": view_type.value,
                    "view_label": VIEW_LABELS[view_type],
                    "image_path": snapshot["export_images"].get(view_type.value),
                }
                for view_type in VIEW_ORDER
            ],
        }

    def build_export_preview(self, session_id: int) -> tuple[dict[str, Any], dict[str, Any]]:
        inspection_session = self.inspection_service.get_session_model(session_id)
        snapshot = self.build_snapshot(inspection_session)
        export_payload = self.build_export_payload(snapshot)
        self.session.commit()
        return snapshot, export_payload

    def _build_preview_text(self, snapshot: dict[str, Any]) -> str:
        vehicle = snapshot["vehicle"]
        lines = [
            "Акт осмотра автомобиля",
            f"Заказ #{snapshot['order']['id']}",
            f"Клиент: {snapshot['client']['full_name']}",
            f"Автомобиль: {' '.join(part for part in [vehicle.get('brand'), vehicle.get('model')] if part)}",
            f"Госномер: {vehicle.get('plate_number') or '—'}",
            f"VIN: {vehicle.get('vin') or '—'}",
            f"Отметок: {len(snapshot['marks'])}",
            f"Общих фото: {len(snapshot['general_photos'])}",
        ]
        if snapshot["session"].get("general_comment"):
            lines.append(f"Комментарий: {snapshot['session']['general_comment']}")
        return "\n".join(lines)

    def _build_template_bundle(self, snapshot: dict[str, Any]) -> TemplateRenderBundle:
        preview_text = self._build_preview_text(snapshot)
        return TemplateRenderBundle(
            template_placeholders=[],
            unresolved_placeholders=[],
            manual_placeholders=[],
            manual_fields_by_design=[],
            common_replacements={},
            service_rows=[],
            material_rows=[],
            preview_content=preview_text,
            section_flags={},
        )

    @staticmethod
    def _safe_text(value: Any, *, fallback: str = "—") -> str:
        if value is None:
            return fallback
        text = str(value).strip()
        if not text:
            return fallback
        return text

    @staticmethod
    def _chunk_items(items: list[dict[str, Any]], size: int) -> list[list[dict[str, Any]]]:
        return [items[index : index + size] for index in range(0, len(items), size)]

    def _add_photo_gallery(self, document, photos: list[dict[str, Any]], *, image_width, columns: int = 2) -> None:
        resolved_photos = []
        for photo in photos:
            file_path = photo.get("file_path")
            if not file_path:
                continue
            resolved_path = self.inspection_service.storage.resolve(file_path)
            if resolved_path.exists():
                resolved_photos.append({"path": resolved_path, "caption": self._safe_text(photo.get("original_name"), fallback="Фотография")})

        if not resolved_photos:
            document.add_paragraph("Фотографии не приложены.")
            return

        for row in self._chunk_items(resolved_photos, columns):
            table = document.add_table(rows=1, cols=columns)
            for index, cell in enumerate(table.rows[0].cells):
                if index >= len(row):
                    cell.text = ""
                    continue
                cell.paragraphs[0].add_run().add_picture(str(row[index]["path"]), width=image_width)
                cell.add_paragraph(row[index]["caption"])

    def _write_docx(self, snapshot: dict[str, Any], output_path: Path) -> None:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except Exception as exc:  # pragma: no cover - runtime dependency
            raise AppError(
                code="docx_generation_unavailable",
                message="Для генерации акта осмотра нужен пакет python-docx",
                status_code=500,
            ) from exc

        document = Document()
        title = document.add_heading("Акт осмотра автомобиля", level=0)
        title.runs[0].font.size = Pt(18)

        order_id = snapshot["order"]["id"]
        vehicle_display = " ".join(
            part for part in [self._safe_text(snapshot["vehicle"].get("brand"), fallback=""), self._safe_text(snapshot["vehicle"].get("model"), fallback="")] if part
        ) or "—"

        document.add_paragraph(f"Заказ: №{order_id}")
        document.add_paragraph(f"Дата формирования: {self._safe_text(snapshot.get('generated_at'), fallback=datetime.now().strftime('%d.%m.%Y %H:%M'))}")
        document.add_paragraph(f"Клиент: {self._safe_text(snapshot['client'].get('full_name'))}")
        document.add_paragraph(f"Телефон: {self._safe_text(snapshot['client'].get('phone'))}")
        document.add_paragraph(f"Автомобиль: {vehicle_display}")
        document.add_paragraph(f"Госномер: {self._safe_text(snapshot['vehicle'].get('plate_number'))}")
        document.add_paragraph(f"VIN: {self._safe_text(snapshot['vehicle'].get('vin'))}")
        document.add_paragraph(f"Осмотр выполнил: {self._safe_text(snapshot['session'].get('inspector_name'))}")

        general_comment = self._safe_text(snapshot["session"].get("general_comment"), fallback="")
        if general_comment:
            document.add_heading("Общий комментарий", level=1)
            document.add_paragraph(general_comment)
        else:
            document.add_heading("Общий комментарий", level=1)
            document.add_paragraph("Комментарий не указан.")

        checklist = snapshot["session"].get("checklist_json") or {}
        document.add_heading("Чек-лист", level=1)
        if checklist:
            table = document.add_table(rows=1, cols=3)
            header_cells = table.rows[0].cells
            header_cells[0].text = "Пункт"
            header_cells[1].text = "Состояние"
            header_cells[2].text = "Комментарий"
            for key, value in checklist.items():
                row_cells = table.add_row().cells
                if isinstance(value, dict):
                    row_cells[0].text = self._safe_text(value.get("label"), fallback=key)
                    row_cells[1].text = "Да" if value.get("checked") else "Нет"
                    row_cells[2].text = self._safe_text(value.get("note"), fallback="")
                else:
                    row_cells[0].text = key
                    row_cells[1].text = "Да" if value else "Нет"
                    row_cells[2].text = ""
        else:
            document.add_paragraph("Чек-лист не заполнен.")

        document.add_heading("Проекции автомобиля", level=1)
        rendered_view_count = 0
        for view in VIEW_ORDER:
            image_path = snapshot["export_images"].get(view.value)
            if not image_path:
                continue
            document.add_paragraph(VIEW_LABELS[view], style="Heading 2")
            resolved_path = self.inspection_service.storage.resolve(image_path)
            if resolved_path.exists():
                document.add_picture(str(resolved_path), width=Inches(6.4))
                rendered_view_count += 1
        if rendered_view_count == 0:
            document.add_paragraph("Проекции с отметками не сформированы.")

        document.add_heading("Список дефектов", level=1)
        if snapshot["marks"]:
            defects_table = document.add_table(rows=1, cols=6)
            header = defects_table.rows[0].cells
            header[0].text = "№"
            header[1].text = "Проекция"
            header[2].text = "Тип"
            header[3].text = "Степень"
            header[4].text = "Статус"
            header[5].text = "Комментарий"
            for item in snapshot["marks"]:
                cells = defects_table.add_row().cells
                cells[0].text = str(item["number"])
                cells[1].text = self._safe_text(item.get("view_label"))
                cells[2].text = self._safe_text(item.get("defect_type"))
                cells[3].text = self._safe_text(item.get("severity"))
                cells[4].text = self._safe_text(item.get("status"))
                cells[5].text = self._safe_text(item.get("comment"), fallback="")
        else:
            document.add_paragraph("Дефекты не зафиксированы.")

        mark_photos = [item for item in snapshot["marks"] if item["photos"]]
        document.add_heading("Фотографии по дефектам", level=1)
        if mark_photos:
            for item in mark_photos:
                document.add_paragraph(
                    f"#{item['number']} · {self._safe_text(item.get('view_label'))} · {self._safe_text(item.get('defect_type'))}",
                    style="Heading 2",
                )
                if item.get("comment"):
                    document.add_paragraph(self._safe_text(item.get("comment"), fallback=""))
                self._add_photo_gallery(document, item["photos"], image_width=Inches(2.7))
        else:
            document.add_paragraph("Фотографии по дефектам отсутствуют.")

        document.add_heading("Общие фотографии", level=1)
        if snapshot["general_photos"]:
            self._add_photo_gallery(document, snapshot["general_photos"], image_width=Inches(2.7))
        else:
            document.add_paragraph("Общие фотографии не приложены.")

        document.add_paragraph("")
        signatures = document.add_table(rows=2, cols=2)
        signatures.rows[0].cells[0].text = "Сотрудник"
        signatures.rows[0].cells[1].text = "Клиент"
        signatures.rows[1].cells[0].text = "________________"
        signatures.rows[1].cells[1].text = "________________"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)

    def generate_act(self, session_id: int, actor_user_id: int) -> tuple[CrmInspectionExport, CrmDocument]:
        inspection_session = self.inspection_service.get_session_model(session_id)
        if inspection_session.status not in {
            InspectionSessionStatus.COMPLETED,
            InspectionSessionStatus.CONFIRMED,
            InspectionSessionStatus.LOCKED,
        }:
            raise AppError(
                code="inspection_not_ready",
                message="Сформировать акт можно только после завершения осмотра",
                status_code=409,
            )

        snapshot = self.build_snapshot(inspection_session)
        export_payload = self.build_export_payload(snapshot)
        document = self._create_document(inspection_session, actor_user_id)
        bundle = self._build_template_bundle(snapshot)
        export_record = CrmInspectionExport(
            inspection_session_id=inspection_session.id,
            document_id=document.id,
            snapshot_json=snapshot,
            export_payload_json=export_payload,
            created_by_user_id=actor_user_id,
        )
        self.session.add(export_record)
        self.session.flush()

        docx_path = self.document_storage.get_docx_output_path(
            document.order_id,
            document.id,
            document.document_type,
            document.document_number,
            document.created_at,
        )
        pdf_path = self.document_storage.get_pdf_output_path(
            document.order_id,
            document.id,
            document.document_type,
            document.document_number,
            document.created_at,
        )
        self._write_docx(snapshot, docx_path)
        conversion = self.pdf_converter.convert(docx_path, pdf_path, bundle)

        document.storage_docx_path = str(docx_path)
        document.storage_pdf_path = str(pdf_path)
        document.last_rendered_at = app_now_naive()
        document.last_pdf_engine = conversion.engine
        document.last_pdf_generation_note = conversion.note
        document.work_started_at = inspection_session.started_at
        document.work_completed_at = inspection_session.completed_at
        document.updated_at = app_now_naive()

        export_record.docx_path = str(docx_path)
        export_record.pdf_path = str(pdf_path)

        self.audit_logs.record(
            actor_user_id=actor_user_id,
            entity_type="inspection_session",
            entity_id=inspection_session.id,
            action="inspection_act_generated",
            title="Сформирован акт осмотра автомобиля",
            description=f"Документ #{document.document_number}",
        )
        self.session.commit()
        return export_record, document

    def build_document_preview(self, document_id: int) -> DocumentPreviewRead:
        export = self._get_export_by_document(document_id)
        snapshot = export.snapshot_json
        return DocumentPreviewRead(
            document_id=document_id,
            document_type=DocumentType.INSPECTION_ACT,
            document_number=self._get_document(document_id).document_number,
            preview_content=self._build_preview_text(snapshot),
            is_placeholder_rendering=False,
            rendered_at=export.created_at,
            unresolved_placeholders=[],
            manual_fields_by_design=[],
        )

    def render_document(self, document_id: int) -> CrmDocument:
        export = self._get_export_by_document(document_id)
        document = self._get_document(document_id)
        docx_path = self.document_storage.get_docx_output_path(
            document.order_id,
            document.id,
            document.document_type,
            document.document_number,
            document.created_at,
        )
        self._write_docx(export.snapshot_json, docx_path)
        document.storage_docx_path = str(docx_path)
        document.last_rendered_at = app_now_naive()
        self.session.commit()
        return document

    def generate_pdf(self, document_id: int) -> CrmDocument:
        export = self._get_export_by_document(document_id)
        document = self.render_document(document_id)
        bundle = self._build_template_bundle(export.snapshot_json)
        docx_path = Path(document.storage_docx_path or "")
        pdf_path = self.document_storage.get_pdf_output_path(
            document.order_id,
            document.id,
            document.document_type,
            document.document_number,
            document.created_at,
        )
        conversion = self.pdf_converter.convert(docx_path, pdf_path, bundle)
        document.storage_pdf_path = str(pdf_path)
        document.last_pdf_engine = conversion.engine
        document.last_pdf_generation_note = conversion.note
        document.last_rendered_at = app_now_naive()
        export.pdf_path = str(pdf_path)
        self.session.commit()
        return document
